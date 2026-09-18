from __future__ import annotations

import csv
from dataclasses import dataclass
from io import StringIO
from pathlib import Path
from typing import Callable, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .domain import ALIASES, REQUIRED_FIELDS, Movement, ParseIssue, normalize_row


@dataclass(frozen=True, slots=True)
class ReadResult:
    movements: list[Movement]
    issues: list[ParseIssue]
    total_rows: int


COLPPY_INSTRUCTIONS = [
    "Campo obligatorio                                               Formato dd-mm-aaaa",
    "Campo obligatorio",
    "",
    "Campo obligatorio                                               Numérico 2 (dos) decimales con valor negativo para los débitos",
]
COLPPY_HEADERS = ["Fecha", "Concepto", "Nro. Comprobante", "Importe"]
ProgressCallback = Callable[[int, str], None]


def _has_required_columns(headers: Sequence[str]) -> bool:
    header_set = {str(header).strip() for header in headers if header is not None}
    return all(any(alias in header_set for alias in ALIASES[field]) for field in REQUIRED_FIELDS)


def _read_csv_rows(path: Path) -> tuple[list[str], list[list[str]]]:
    raw = path.read_bytes()
    text: str | None = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("El CSV no usa una codificación compatible")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=";,\t")
        reader = csv.reader(StringIO(text), dialect)
        rows = list(reader)
    except csv.Error as exc:
        raise ValueError(f"No se pudo reconocer el separador del CSV: {exc}") from exc
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _read_xls_rows(path: Path) -> tuple[list[object], list[list[object]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("Falta la dependencia xlrd para leer archivos .xls") from exc
    workbook = xlrd.open_workbook(filename=str(path), on_demand=True)
    try:
        if workbook.nsheets == 0:
            return [], []
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows == 0:
            return [], []
        headers = list(sheet.row_values(0))
        rows: list[list[object]] = []
        for row_index in range(1, sheet.nrows):
            values: list[object] = []
            for cell in sheet.row(row_index):
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                values.append(value)
            rows.append(values)
        return headers, rows
    finally:
        workbook.release_resources()


def _consume_rows(
    path: Path,
    raw_headers: Sequence[object],
    rows: Iterable[Sequence[object]],
    estimated_rows: int,
    movements: list[Movement],
    issues: list[ParseIssue],
    total_rows: int,
    file_start: int,
    file_end: int,
    report: ProgressCallback,
) -> int:
    headers = ["" if value is None else str(value).strip() for value in raw_headers]
    if not _has_required_columns(headers):
        issues.append(
            ParseIssue(
                path.name,
                1,
                "Faltan columnas obligatorias de Mercado Pago: ID, tipo, fecha o importe neto",
            )
        )
        report(file_end, f"Archivo revisado: {path.name}")
        return total_rows

    for row_number, values in enumerate(rows, start=2):
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
            continue
        total_rows += 1
        row = dict(zip(headers, values))
        try:
            movements.append(normalize_row(row, path.name, row_number))
        except ValueError as exc:
            issues.append(ParseIssue(path.name, row_number, str(exc)))
        processed_rows = row_number - 1
        if processed_rows == 1 or processed_rows % 250 == 0 or processed_rows >= estimated_rows:
            fraction = min(1.0, processed_rows / max(1, estimated_rows))
            value = file_start + int((file_end - file_start) * fraction)
            report(value, f"Leyendo {path.name}: {processed_rows} filas")
    return total_rows


def read_mercadopago_files(
    paths: Iterable[str | Path],
    progress: ProgressCallback | None = None,
) -> ReadResult:
    movements: list[Movement] = []
    issues: list[ParseIssue] = []
    total_rows = 0
    path_list = [Path(path) for path in paths]

    def report(value: int, message: str) -> None:
        if progress is not None:
            progress(max(0, min(100, int(value))), message)

    report(0, "Preparando archivos...")
    if not path_list:
        report(100, "Análisis completado")
        return ReadResult(movements, issues, total_rows)

    file_count = len(path_list)
    for file_index, path in enumerate(path_list):
        file_start = int(file_index * 100 / file_count)
        file_end = int((file_index + 1) * 100 / file_count)
        report(file_start, f"Abriendo {path.name}...")
        if path.suffix.lower() in {".csv", ".xls"}:
            try:
                if path.suffix.lower() == ".csv":
                    raw_headers, tabular_rows = _read_csv_rows(path)
                else:
                    raw_headers, tabular_rows = _read_xls_rows(path)
                if not raw_headers:
                    issues.append(ParseIssue(path.name, 0, "El archivo está vacío"))
                else:
                    total_rows = _consume_rows(
                        path,
                        raw_headers,
                        tabular_rows,
                        max(1, len(tabular_rows)),
                        movements,
                        issues,
                        total_rows,
                        file_start,
                        file_end,
                        report,
                    )
            except Exception as exc:
                issues.append(ParseIssue(path.name, 0, f"No se pudo abrir el archivo: {exc}"))
            report(file_end, f"Archivo revisado: {path.name}")
            continue
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            issues.append(ParseIssue(path.name, 0, f"No se pudo abrir el archivo: {exc}"))
            report(file_end, f"Archivo revisado: {path.name}")
            continue

        try:
            sheet = workbook.worksheets[0]
            rows = sheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration:
                issues.append(ParseIssue(path.name, 0, "El archivo está vacío"))
                continue
            estimated_rows = max(1, int(sheet.max_row or 1) - 1)
            total_rows = _consume_rows(
                path,
                raw_headers,
                rows,
                estimated_rows,
                movements,
                issues,
                total_rows,
                file_start,
                file_end,
                report,
            )
        finally:
            workbook.close()
        report(file_end, f"Archivo revisado: {path.name}")

    report(100, "Análisis completado")
    return ReadResult(movements, issues, total_rows)


def _sorted_movements(movements: Iterable[Movement]) -> list[Movement]:
    return sorted(movements, key=lambda item: (item.date, item.occurred_at, item.receipt))


def export_colppy_csv(movements: Iterable[Movement], output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="cp1252", newline="") as stream:
        writer = csv.writer(stream, delimiter=";", lineterminator="\r\n")
        writer.writerow(COLPPY_INSTRUCTIONS)
        writer.writerow(COLPPY_HEADERS)
        for item in _sorted_movements(movements):
            date_text = f"{item.date.day}/{item.date.month}/{item.date.year}"
            amount_text = format(item.amount, ".2f").replace(".", ",")
            writer.writerow([date_text, item.concept, item.receipt, amount_text])
    return output


def export_colppy_xlsx(movements: Iterable[Movement], output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clientes"
    sheet.append(COLPPY_INSTRUCTIONS)
    sheet.append(COLPPY_HEADERS)

    for item in _sorted_movements(movements):
        sheet.append([item.date, item.concept, item.receipt, float(item.amount)])

    instruction_fill = PatternFill("solid", fgColor="E7E6E6")
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=9, italic=True, color="595959")
        cell.fill = instruction_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for cell in sheet[2]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(3, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "dd-mm-yyyy"
        sheet.cell(row, 4).number_format = "#,##0.00;[Red]-#,##0.00"
        for col in range(1, 5):
            sheet.cell(row, col).font = Font(name="Arial", size=10)
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 62
    sheet.column_dimensions["C"].width = 24
    sheet.column_dimensions["D"].width = 20
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A3"
    workbook.save(output)
    workbook.close()
    return output


def export_colppy_xls(movements: Iterable[Movement], output_path: str | Path) -> Path:
    try:
        import xlwt
    except ImportError as exc:
        raise RuntimeError("Falta la dependencia xlwt para generar archivos .xls") from exc

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet("Clientes")

    instructions = xlwt.easyxf("font: name Arial, height 180, italic on, colour gray40; pattern: pattern solid, fore_colour gray25; alignment: wrap on, vert centre")
    headers = xlwt.easyxf("font: name Arial, height 200, bold on, colour white; pattern: pattern solid, fore_colour dark_blue; alignment: horiz center, vert centre")
    date_style = xlwt.easyxf("font: name Arial, height 200", num_format_str="DD-MM-YYYY")
    text_style = xlwt.easyxf("font: name Arial, height 200")
    amount_style = xlwt.easyxf("font: name Arial, height 200", num_format_str="#,##0.00;[Red]-#,##0.00")

    for col, value in enumerate(COLPPY_INSTRUCTIONS):
        sheet.write(0, col, value, instructions)
    for col, value in enumerate(COLPPY_HEADERS):
        sheet.write(1, col, value, headers)
    for row, item in enumerate(_sorted_movements(movements), start=2):
        sheet.write(row, 0, item.date, date_style)
        sheet.write(row, 1, item.concept, text_style)
        sheet.write(row, 2, item.receipt, text_style)
        sheet.write(row, 3, float(item.amount), amount_style)

    sheet.col(0).width = 14 * 256
    sheet.col(1).width = 62 * 256
    sheet.col(2).width = 24 * 256
    sheet.col(3).width = 20 * 256
    sheet.row(0).height_mismatch = True
    sheet.row(0).height = 32 * 20
    workbook.save(str(output))
    return output
