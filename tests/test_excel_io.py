from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import importlib.util
import csv
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from src.mp_colppy.domain import normalize_row
from src.mp_colppy.excel_io import export_colppy_csv, export_colppy_xls, export_colppy_xlsx, read_mercadopago_files


HEADERS = [
    "ID DE OPERACIÓN EN MERCADO PAGO",
    "TIPO DE OPERACIÓN",
    "FECHA DE ORIGEN",
    "MONTO NETO DE LA OPERACIÓN QUE IMPACTÓ TU DINERO",
    "PAGADOR",
    "MEDIO DE PAGO",
    "NOMBRE DE LOCAL",
]


class ExcelReadTests(unittest.TestCase):
    def test_reads_excel_date_cell_from_legacy_xls_report(self):
        import xlwt

        with TemporaryDirectory() as temp:
            path = Path(temp) / "reporte-con-fecha.xls"
            workbook = xlwt.Workbook()
            sheet = workbook.add_sheet("Reporte")
            for column, header in enumerate(HEADERS):
                sheet.write(0, column, header)
            date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY HH:MM:SS")
            values = [
                303,
                "Pago aprobado",
                datetime(2026, 9, 3, 12, 45),
                975.25,
                "Cliente de prueba",
                "QR",
                "Local de prueba",
            ]
            for column, value in enumerate(values):
                style = date_style if column == 2 else xlwt.Style.default_style
                sheet.write(1, column, value, style)
            workbook.save(str(path))

            result = read_mercadopago_files([path])

        self.assertEqual(result.total_rows, 1)
        self.assertEqual(result.issues, [])
        self.assertEqual(result.movements[0].date.isoformat(), "2026-09-03")

    def test_reads_legacy_xls_report(self):
        import xlwt

        with TemporaryDirectory() as temp:
            path = Path(temp) / "reporte.xls"
            workbook = xlwt.Workbook()
            sheet = workbook.add_sheet("Reporte")
            for column, header in enumerate(HEADERS):
                sheet.write(0, column, header)
            values = [
                202,
                "Pago aprobado",
                "2026-09-02T11:30:00-03:00",
                2500.75,
                "Cliente de prueba",
                "QR",
                "Local de prueba",
            ]
            for column, value in enumerate(values):
                sheet.write(1, column, value)
            workbook.save(str(path))

            result = read_mercadopago_files([path])

        self.assertEqual(result.total_rows, 1)
        self.assertEqual(len(result.issues), 0)
        self.assertEqual(result.movements[0].amount, Decimal("2500.75"))
        self.assertEqual(result.movements[0].receipt, "202")

    def test_reads_semicolon_csv_with_decimal_comma(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / "reporte.csv"
            content = (
                ";".join(HEADERS)
                + "\r\n"
                + ";".join(
                    [
                        "101",
                        "Pago aprobado",
                        "2026-09-01T10:00:00-03:00",
                        "1234,50",
                        "Cliente de prueba",
                        "QR",
                        "Local de prueba",
                    ]
                )
                + "\r\n"
            )
            path.write_bytes(content.encode("cp1252"))

            result = read_mercadopago_files([path])

        self.assertEqual(result.total_rows, 1)
        self.assertEqual(len(result.issues), 0)
        self.assertEqual(result.movements[0].amount, Decimal("1234.50"))
        self.assertEqual(result.movements[0].receipt, "101")

    def test_reads_comma_delimited_csv(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / "reporte-comas.csv"
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream, delimiter=",")
                writer.writerow(HEADERS)
                writer.writerow(
                    [
                        "102",
                        "Pago aprobado",
                        "2026-09-02T10:00:00-03:00",
                        "987.65",
                        "Cliente, con coma",
                        "QR",
                        "Local de prueba",
                    ]
                )

            result = read_mercadopago_files([path])

        self.assertEqual(result.total_rows, 1)
        self.assertEqual(result.issues, [])
        self.assertEqual(result.movements[0].amount, Decimal("987.65"))
        self.assertIn("Cliente, con coma", result.movements[0].concept)

    def test_reads_valid_rows_and_collects_invalid_rows(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / "mp.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(HEADERS)
            sheet.append([101, "Pago aprobado", "2026-09-01T10:00:00-03:00", 1000, "Cliente A", "QR", "Local de prueba"])
            sheet.append([None, "Pago aprobado", "2026-09-01T11:00:00-03:00", 2000, "Cliente B", "QR", "Local de prueba"])
            workbook.save(path)

            result = read_mercadopago_files([path])

        self.assertEqual(result.total_rows, 2)
        self.assertEqual(len(result.movements), 1)
        self.assertEqual(len(result.issues), 1)
        self.assertEqual(result.movements[0].receipt, "101")
        self.assertIn("ID de operación", result.issues[0].message)

    def test_rejects_workbook_without_required_columns(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / "otro.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Fecha", "Importe"])
            sheet.append(["2026-09-01", 10])
            workbook.save(path)

            result = read_mercadopago_files([path])

        self.assertEqual(len(result.movements), 0)
        self.assertEqual(len(result.issues), 1)
        self.assertIn("columnas obligatorias", result.issues[0].message)


class ExcelExportTests(unittest.TestCase):
    def test_exports_official_colppy_csv_bytes(self):
        item = normalize_row(
            {
                "SOURCE_ID": 17,
                "TRANSACTION_TYPE": "Pago aprobado",
                "TRANSACTION_DATE": "2026-09-01T10:00:00-03:00",
                "SETTLEMENT_NET_AMOUNT": "1234.50",
                "PAYER_NAME": "Cliente de prueba",
            },
            "origen.xlsx",
            2,
        )
        expected = (
            "Campo obligatorio                                               Formato dd-mm-aaaa;"
            "Campo obligatorio;;"
            "Campo obligatorio                                               Numérico 2 (dos) decimales con valor negativo para los débitos\r\n"
            "Fecha;Concepto;Nro. Comprobante;Importe\r\n"
            "1/9/2026;Pago aprobado - Cliente de prueba;17;1234,50\r\n"
        ).encode("cp1252")

        with TemporaryDirectory() as temp:
            path = Path(temp) / "colppy.csv"
            export_colppy_csv([item], path)

            self.assertEqual(path.read_bytes(), expected)

    def test_exports_two_colppy_header_rows_and_sorted_movements(self):
        later = normalize_row(
            {
                "SOURCE_ID": 2,
                "TRANSACTION_TYPE": "PAYOUTS",
                "TRANSACTION_DATE": "2026-09-02T10:00:00-03:00",
                "SETTLEMENT_NET_AMOUNT": -500,
                "PAYMENT_METHOD": "Transferencia",
            },
            "dos.xlsx",
            2,
        )
        earlier = normalize_row(
            {
                "SOURCE_ID": 1,
                "TRANSACTION_TYPE": "Pago aprobado",
                "TRANSACTION_DATE": "2026-09-01T10:00:00-03:00",
                "SETTLEMENT_NET_AMOUNT": 1000.25,
                "PAYER_NAME": "Ana",
            },
            "uno.xlsx",
            2,
        )

        with TemporaryDirectory() as temp:
            path = Path(temp) / "colppy.xlsx"
            export_colppy_xlsx([later, earlier], path)
            workbook = load_workbook(path, data_only=True)
            sheet = workbook.active

            self.assertEqual(sheet["A1"].value, "Campo obligatorio                                               Formato dd-mm-aaaa")
            self.assertEqual([sheet.cell(2, col).value for col in range(1, 5)], ["Fecha", "Concepto", "Nro. Comprobante", "Importe"])
            self.assertEqual(sheet["C3"].value, "1")
            self.assertEqual(sheet["D3"].value, 1000.25)
            self.assertEqual(sheet["C4"].value, "2-PAYOUT")
            self.assertEqual(sheet["D4"].value, -500)
            self.assertEqual(sheet["A3"].number_format, "dd-mm-yyyy")
            workbook.close()

    def test_xls_export_is_available_or_reports_missing_dependency(self):
        item = normalize_row(
            {
                "SOURCE_ID": 1,
                "TRANSACTION_TYPE": "Pago aprobado",
                "TRANSACTION_DATE": "2026-09-01",
                "SETTLEMENT_NET_AMOUNT": 10,
            },
            "uno.xlsx",
            2,
        )
        with TemporaryDirectory() as temp:
            path = Path(temp) / "colppy.xls"
            if importlib.util.find_spec("xlwt") is None:
                with self.assertRaisesRegex(RuntimeError, "xlwt"):
                    export_colppy_xls([item], path)
            else:
                export_colppy_xls([item], path)
                self.assertGreater(path.stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
